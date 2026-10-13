"""BMCC Advising Agent.

A command-line advising assistant for Blue Mountain Community College staff.
It reads the advising library and exported degree audits, looks up the BMCC
catalog and other colleges' programs on the web, and fills out term-by-term
degree planning worksheets, dual credit included.

Folders (all next to this file):
  library/     advising guides, worksheets, catalog PDFs staff drop in
  audits/      degree audit exports from the SIS (PDF, DOCX or text)
  worksheets/  completed planning worksheets the agent writes
"""

import datetime
import pathlib
import re
import sys
from typing import List, Optional

from typing_extensions import TypedDict

import anthropic
import docx
import pypdf
from anthropic import beta_tool

MODEL = "claude-opus-5-5"
HERE = pathlib.Path(__file__).parent.resolve()
LIBRARY = HERE / "library"
AUDITS = HERE / "audits"
WORKSHEETS = HERE / "worksheets"
for folder in (LIBRARY, AUDITS, WORKSHEETS):
    folder.mkdir(exist_ok=True)

READABLE = {".pdf", ".docx", ".txt", ".md", ".csv"}
MAX_TERM_CREDITS = 18

SYSTEM = """You are an academic advising assistant for Blue Mountain Community College (BMCC) in eastern Oregon, working with BMCC advisors and staff. You help build complete, realistic pathways for students: BMCC certificates and degrees, dual credit earned in high school, and transfer to four-year colleges and universities.

Sources, in order of authority:
1. The student's degree audit (read_audit). It is the official record of what the student has completed.
2. The BMCC catalog for the student's catalog year, at https://catalog.bluecc.edu/ (Modern Campus catalog; each program page has a Degree Planner checklist).
3. BMCC advising material in the local library (list_library, read_library_document), and on the web at https://bluecc.edu/advising/ and https://www.bluecc.edu/about-bmcc/staff-web/advising-resources
4. Dual credit, Sponsored Dual Credit and Expanded Options: https://bluecc.edu/academics/early-college-credit
5. Other institutions' official catalogs and transfer pages, for transfer pathways. Prefer Oregon public universities and the Oregon transfer frameworks (AAOT, ASOT, Major Transfer Maps) when relevant, and check each receiving school's own published transfer rules.

How to work:
- Check every course number, credit value and requirement against a source you actually read in this conversation. Never invent a course, prerequisite, credit value or transfer equivalency. If you cannot verify something, say so plainly and mark it "verify with advisor" on the worksheet.
- Respect prerequisites and sequencing (for example, math and writing sequences) and typical term offerings when the catalog lists them.
- Count dual credit and transfer credit already on the audit before planning new courses, and do not schedule a course the student has completed.
- Keep term loads realistic: full time is 12 to 18 credits unless the advisor says otherwise.
- When the plan is ready, call save_worksheet. Include the source URLs or file names you relied on.
- Degree audits contain student records protected by FERPA. Use them only for the advising task in front of you and do not repeat identifying details beyond what the worksheet needs.
- Write for an advisor: short, specific, no filler. Flag risks (missing prerequisites, courses offered only once a year, transfer requirements BMCC does not cover).
"""


# ---------- reading files ----------

def _extract_text(path: pathlib.Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = pypdf.PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="replace")


def _safe_path(root: pathlib.Path, relative: str) -> pathlib.Path:
    path = (root / relative).resolve()
    if not path.is_relative_to(root):
        raise ValueError(f"{relative!r} is outside {root.name}/")
    return path


def _list(root: pathlib.Path) -> str:
    files = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in READABLE)
    if not files:
        return f"No readable files in {root.name}/ ({', '.join(sorted(READABLE))})."
    return "\n".join(str(p.relative_to(root)) for p in files)


def _read(root: pathlib.Path, name: str) -> str:
    try:
        path = _safe_path(root, name)
        if path.suffix.lower() not in READABLE:
            return f"Error: unsupported file type {path.suffix}"
        text = _extract_text(path).strip()
        return text or "Error: no text could be extracted (the file may be a scanned image)."
    except (OSError, ValueError, pypdf.errors.PdfReadError) as e:
        return f"Error: {e}"


# Rough redaction of identifiers in audit text. It catches SSN-shaped numbers,
# long ID numbers, emails and dates of birth labels; it is not a guarantee.
_REDACTIONS = [
    (re.compile(r"\b\d{3}-\d{2}-\d{4}\b"), "[SSN]"),
    (re.compile(r"\b\d{7,10}\b"), "[ID]"),
    (re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+"), "[EMAIL]"),
    (re.compile(r"(?i)(birth\s*date|date\s*of\s*birth|dob)\s*[:#]?\s*\S+"), r"\1: [DOB]"),
]


def _redact(text: str) -> str:
    for pattern, replacement in _REDACTIONS:
        text = pattern.sub(replacement, text)
    return text


@beta_tool
def list_library() -> str:
    """List BMCC advising documents in the local library (advising guides, degree planning worksheets, catalog excerpts, transfer maps)."""
    return _list(LIBRARY)


@beta_tool
def read_library_document(name: str) -> str:
    """Read the text of one document from the advising library.

    Args:
        name: File name as shown by list_library, e.g. guides/nursing-advising-guide.pdf.
    """
    return _read(LIBRARY, name)


@beta_tool
def list_audits() -> str:
    """List the degree audit exports available in the audits folder."""
    return _list(AUDITS)


@beta_tool
def read_audit(name: str) -> str:
    """Read a student's degree audit export (completed, in-progress and missing requirements, including dual credit and transfer credit). Identifier numbers and emails are masked.

    Args:
        name: File name as shown by list_audits.
    """
    text = _read(AUDITS, name)
    return text if text.startswith("Error") else _redact(text)


# ---------- writing the worksheet ----------

class Course(TypedDict):
    code: str
    title: str
    credits: float
    requirement: str
    note: str


class Term(TypedDict):
    term: str
    courses: List[Course]


class CompletedCredit(TypedDict):
    code: str
    title: str
    credits: float
    source: str


def _validate(terms: List[Term], completed: List[CompletedCredit], total_required: float) -> List[str]:
    warnings = []
    seen = {c["code"].upper().replace(" ", "") for c in completed}
    for term in terms:
        load = sum(c["credits"] for c in term["courses"])
        if load > MAX_TERM_CREDITS:
            warnings.append(f"{term['term']}: {load:g} credits is above {MAX_TERM_CREDITS}.")
        for course in term["courses"]:
            key = course["code"].upper().replace(" ", "")
            if key in seen:
                warnings.append(f"{course['code']} appears more than once (completed or already planned).")
            seen.add(key)
    planned = sum(c["credits"] for t in terms for c in t["courses"])
    done = sum(c["credits"] for c in completed)
    if total_required and planned + done < total_required:
        warnings.append(f"Plan reaches {planned + done:g} of {total_required:g} required credits.")
    return warnings


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "plan"


def _write_docx(path: pathlib.Path, header: dict, completed: List[CompletedCredit],
                terms: List[Term], warnings: List[str], sources: List[str], notes: str) -> None:
    document = docx.Document()
    document.add_heading("Blue Mountain Community College", level=0)
    document.add_heading("Degree Planning Worksheet", level=1)
    for label, value in header.items():
        if value:
            p = document.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(str(value))

    if completed:
        document.add_heading("Credit already earned (dual credit, transfer, prior BMCC)", level=2)
        table = document.add_table(rows=1, cols=4, style="Light Grid Accent 1")
        for cell, text in zip(table.rows[0].cells, ("Course", "Title", "Credits", "Source")):
            cell.text = text
        for c in completed:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text, row[3].text = c["code"], c["title"], f"{c['credits']:g}", c["source"]

    for term in terms:
        load = sum(c["credits"] for c in term["courses"])
        document.add_heading(f"{term['term']} ({load:g} credits)", level=2)
        table = document.add_table(rows=1, cols=5, style="Light Grid Accent 1")
        for cell, text in zip(table.rows[0].cells, ("Course", "Title", "Credits", "Requirement", "Note")):
            cell.text = text
        for c in term["courses"]:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = c["code"], c["title"], f"{c['credits']:g}"
            row[3].text, row[4].text = c["requirement"], c["note"]

    if warnings:
        document.add_heading("Checks to review", level=2)
        for w in warnings:
            document.add_paragraph(w, style="List Bullet")
    if notes:
        document.add_heading("Advisor notes", level=2)
        document.add_paragraph(notes)
    if sources:
        document.add_heading("Sources", level=2)
        for s in sources:
            document.add_paragraph(s, style="List Bullet")
    document.add_paragraph().add_run(
        "Planning tool only. The official degree audit and the catalog for the student's "
        "catalog year govern graduation requirements."
    ).italic = True
    document.save(str(path))


@beta_tool
def save_worksheet(
    student_label: str,
    program: str,
    catalog_year: str,
    total_credits_required: float,
    completed: List[CompletedCredit],
    terms: List[Term],
    transfer_goal: str = "",
    notes: str = "",
    sources: Optional[List[str]] = None,
) -> str:
    """Fill out and save a term-by-term degree planning worksheet as a Word document. Checks term loads, duplicate courses and total credits, and returns any problems found.

    Args:
        student_label: How the advisor refers to the student, e.g. initials or "Student A". Do not use ID numbers.
        program: BMCC certificate or degree, e.g. "Associate of Arts Oregon Transfer (AAOT)".
        catalog_year: Catalog year governing the requirements, e.g. "2026-2027".
        total_credits_required: Total credits the BMCC award requires (from the catalog).
        completed: Credit already earned and applied, including dual credit and transfer credit. Source is e.g. "Dual credit, Hermiston HS" or "BMCC Fall 2025".
        terms: Planned terms in order. Term is e.g. "Fall 2026". Requirement is the requirement each course fills; note is optional (prerequisite, offered fall only, verify with advisor).
        transfer_goal: Receiving institution and major, if the student plans to transfer.
        notes: Short advisor-facing notes: risks, decisions the student needs to make, next steps.
        sources: URLs or library file names the plan relied on.
    """
    sources = sources or []
    warnings = _validate(terms, completed, total_credits_required)
    stamp = datetime.date.today().isoformat()
    path = WORKSHEETS / f"{stamp}-{_slug(student_label)}-{_slug(program)}.docx"
    header = {
        "Student": student_label,
        "Program": program,
        "Catalog year": catalog_year,
        "Transfer goal": transfer_goal,
        "Credits required": f"{total_credits_required:g}",
        "Prepared": stamp,
    }
    try:
        _write_docx(path, header, completed, terms, warnings, sources, notes)
    except OSError as e:
        return f"Error saving worksheet: {e}"
    result = f"Saved {path.relative_to(HERE)}."
    if warnings:
        result += " Checks to review:\n- " + "\n- ".join(warnings)
    return result


# ---------- the agent loop ----------

TOOLS = [
    list_library,
    read_library_document,
    list_audits,
    read_audit,
    save_worksheet,
    {"type": "web_search_20260209", "name": "web_search", "max_uses": 10},
    {"type": "web_fetch_20260209", "name": "web_fetch", "max_uses": 15},
]

MAX_PAUSE_RESTARTS = 5


def run_turn(client: anthropic.Anthropic, messages: list) -> None:
    """Run one advisor request to completion, appending everything to `messages`."""
    for _ in range(MAX_PAUSE_RESTARTS + 1):
        runner = client.beta.messages.tool_runner(
            model=MODEL,
            max_tokens=16000,
            system=SYSTEM,
            tools=TOOLS,
            messages=messages,
            output_config={"effort": "high"},
            # If a safety classifier declines, retry on a suitable fallback model.
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",
        )
        last = None
        for message in runner:
            last = message
            messages.append({"role": "assistant", "content": message.content})
            for block in message.content:
                if block.type == "text" and block.text.strip():
                    print(f"\nAdvisor agent: {block.text}")
                elif block.type == "tool_use":
                    print(f"  [tool] {block.name}")
                elif block.type == "server_tool_use":
                    target = block.input.get("query") or block.input.get("url") or ""
                    print(f"  [{block.name}] {target}")
            tool_response = runner.generate_tool_call_response()
            if tool_response is not None:
                messages.append(tool_response)
        if last is not None and last.stop_reason == "refusal":
            print("\nThe request was declined.")
        # Long web research can pause mid-turn; resume it with a fresh runner.
        if last is None or last.stop_reason != "pause_turn":
            return
    print("\nStopped: the research turn kept pausing. Ask again to continue.")


def main() -> None:
    client = anthropic.Anthropic()
    messages: list = []
    print(
        "BMCC Advising Agent\n"
        f"Library: {LIBRARY}\nAudits: {AUDITS}\nWorksheets: {WORKSHEETS}\n"
        "Example: Build a pathway for audits/student-a.pdf toward the AAOT, then transfer to EOU for business.\n"
        "Type 'exit' to quit."
    )
    while True:
        try:
            user_input = input("\nYou: ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if user_input.lower() in {"exit", "quit"}:
            break
        if not user_input:
            continue
        checkpoint = len(messages)
        messages.append({"role": "user", "content": user_input})
        try:
            run_turn(client, messages)
        except anthropic.AuthenticationError:
            sys.exit("No valid API key. Set ANTHROPIC_API_KEY or run `ant auth login`.")
        except anthropic.RateLimitError:
            print("Rate limited, wait a moment and try again.")
            del messages[checkpoint:]
        except anthropic.APIStatusError as e:
            print(f"API error {e.status_code}: {e.message}")
            del messages[checkpoint:]
        except anthropic.APIConnectionError:
            print("Could not reach the API, check your connection.")
            del messages[checkpoint:]


if __name__ == "__main__":
    main()
